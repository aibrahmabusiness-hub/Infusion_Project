import openpyxl

file_path = r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Step1_Pyxis_Analysis.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb["Order Reconciliation Summary"]

print("Row 2 (First Data Row) Status Columns Details:")
for col_idx in [12, 13, 14, 15]:
    cell = ws.cell(row=2, column=col_idx)
    fill_color = cell.fill.start_color.rgb if cell.fill else None
    font_color = cell.font.color.rgb if cell.font and cell.font.color else None
    print(f"  Col {col_idx} ({ws.cell(row=1, column=col_idx).value}):")
    print(f"    Value: {cell.value}")
    print(f"    Fill Color: {fill_color}")
    print(f"    Font Color: {font_color}")
