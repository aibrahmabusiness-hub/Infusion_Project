import openpyxl

file_path = r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Step1_Pyxis_Analysis.xlsx"
wb = openpyxl.load_workbook(file_path)

if "Order Reconciliation Summary" in wb.sheetnames:
    ws = wb["Order Reconciliation Summary"]
    print("=== Verification of Order Status Colors ===")
    
    count_comp = 0
    count_hang = 0
    
    for r in range(2, ws.max_row + 1):
        order_id = str(ws.cell(row=r, column=1).value)
        status_cell = ws.cell(row=r, column=10)
        status_val = status_cell.value
        
        if status_val == "Completed" and count_comp < 3:
            print(f"Row {r} (Order {order_id}) - Status: {status_val}:")
            print(f"  Fill color (Hex): {status_cell.fill.start_color.rgb if status_cell.fill else 'None'}")
            print(f"  Font color (Hex): {status_cell.font.color.rgb if status_cell.font and status_cell.font.color else 'None'}")
            count_comp += 1
            
        if status_val == "Completed - Max Hang Time" and count_hang < 3:
            print(f"Row {r} (Order {order_id}) - Status: {status_val}:")
            print(f"  Fill color (Hex): {status_cell.fill.start_color.rgb if status_cell.fill else 'None'}")
            print(f"  Font color (Hex): {status_cell.font.color.rgb if status_cell.font and status_cell.font.color else 'None'}")
            count_hang += 1
            
        if count_comp >= 3 and count_hang >= 3:
            break
else:
    print("Sheet 'Order Reconciliation Summary' not found!")
