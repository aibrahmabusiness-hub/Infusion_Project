import docx
from docx.shared import Pt
import os

def create_sdd(output_path):
    doc = docx.Document()
    
    # Title
    title = doc.add_heading('SOLUTION DESIGN DOCUMENT', 0)
    
    # Meta
    doc.add_paragraph('Process Name: Infusion Waste Monitoring')
    doc.add_paragraph('Date: May 12, 2026')
    doc.add_paragraph('Version: V2.0 (Updated from Source Code)')
    
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "This document details the solution design for automating the reconciliation of infusion fluid dispensing versus patient flowsheet charting. "
        "The solution normalizes raw medication data from automated dispensing machines (Pyxis), pairs dispensing (Vends) with corresponding "
        "disposals (Wastes), and compares the computed fluid intake values against charted values in the EMR Flowsheet and Epic Administration records."
    )
    
    doc.add_heading('2. Technical Specifications & Formulas', level=1)
    
    doc.add_heading('2.1 Process Inputs', level=2)
    p_inputs = doc.add_paragraph()
    p_inputs.add_run("The Intelligent Automation bot retrieves the following data sources:\n")
    p_inputs.add_run("1. Pyxis All Device Event Report: ").bold = True
    p_inputs.add_run("Contains medication dispensing and waste logs.\n")
    p_inputs.add_run("2. Epic Administrations Report: ").bold = True
    p_inputs.add_run("Contains nurse medication administration events.\n")
    p_inputs.add_run("3. Epic Flowsheet Report: ").bold = True
    p_inputs.add_run("Contains continuous infusion values logged against the patient.\n")
    p_inputs.add_run("4. Historical State (JSON): ").bold = True
    p_inputs.add_run("A persistent state file to track active carry-over orders across processing weeks.")
    
    doc.add_heading('2.2 Process Steps', level=2)
    
    steps = [
        "Load the consolidated Pyxis, Epic Administrations, and Epic Flowsheet files.",
        "Filter Pyxis, Epic Administrations, and Epic Flowsheet records to include strictly Controlled Substances (e.g., fentanyl, midazolam, versed, ketamine, morphine, hydromorphone, dilaudid).",
        "Normalize patient names (alphanumeric only) and Order IDs to ensure clean matching across datasets.",
        "In the Pyxis dataset, isolate the following Transaction Types: 'Vend', 'Waste', 'Stock Return', and 'Vend (Cancel)'.",
        "Parse the 'MedDescription' column in Pyxis to identify medication name, fluid quantity, and concentration (e.g., 'Total mg = Concentration * Volume').",
        "Execute Vend/Waste Pairing in Pyxis: Group transactions by Order ID and Medication. Chronologically pair 'Waste' transactions with preceding 'Vend' transactions that occurred within a defined hang-time window. Exclude 'Stock Return' and 'Vend (Cancel)' from active pairing.",
        "Calculate the actual Pyxis Infusion volume: \nRemaining Medication (Administered) = Paired Vend Volume - Paired Waste Volume",
        "Match the unique Order IDs across Pyxis, Epic Administrations, and Epic Flowsheet.",
        "Retrieve corresponding total infused values from the Flowsheet for matched orders.",
        "Calculate the Final Variance (Waste Discrepancy): \nFinal Variance = Pyxis Administered Volume - Flowsheet Total Value (per order)",
        "Generate a final consolidated Excel report highlighting exceptions, unmatched items, and order-level variance.",
        "Email the final reconciliation report to the designated stakeholders."
    ]
    
    for idx, step in enumerate(steps, 1):
        doc.add_paragraph(f"{idx}. {step}")
        
    doc.save(output_path)
    print(f"Successfully generated new SDD at {output_path}")

if __name__ == '__main__':
    out_path = r"C:\Users\GaneshBhat\Downloads\Infusion Waste Monitoring - SDD_V2.docx"
    create_sdd(out_path)
