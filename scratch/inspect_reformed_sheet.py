import openpyxl

file_path = r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Step1_Pyxis_Analysis.xlsx"
wb = openpyxl.load_workbook(file_path)

if "Order Reconciliation Summary" in wb.sheetnames:
    ws = wb["Order Reconciliation Summary"]
    print("=== Columns in Sheet 2: Order Reconciliation Summary ===")
    
    headers = [cell.value for cell in ws[1]]
    print("Headers (Length: {}):".format(len(headers)))
    for idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx)
        comment_str = f" | Comment: {cell.comment.text}" if cell.comment else ""
        print(f"  Col {idx}: {h}{comment_str}")
        
    first_row = [cell.value for cell in ws[2]]
    print("\nFirst Row Data:")
    for idx, (h, val) in enumerate(zip(headers, first_row), 1):
        cell = ws.cell(row=2, column=idx)
        print(f"  Col {idx} ({h}): {val} | Format: {cell.number_format}")
else:
    print("Sheet 'Order Reconciliation Summary' not found!")
