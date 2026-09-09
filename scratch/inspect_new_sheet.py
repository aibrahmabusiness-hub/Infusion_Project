import openpyxl

for name, path in [
    ("Step1_Pyxis_Analysis", r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Step1_Pyxis_Analysis.xlsx"),
    ("Infusion_Reconciliation_Report", r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Infusion_Reconciliation_Report.xlsx")
]:
    wb = openpyxl.load_workbook(path)
    print(f"\n=== File: {name} ===")
    if "Administered Calculation" in wb.sheetnames:
        ws = wb["Administered Calculation"]
        headers = [cell.value for cell in ws[1]]
        print(f"  Headers: {headers}")
        
        # Print comments
        print("  Header Comments:")
        for idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=idx)
            if cell.comment:
                print(f"    Col {idx} ({h}): {cell.comment.text}")
                
        first_row = [cell.value for cell in ws[2]]
        print(f"  First Row Data: {first_row}")
    else:
        print("  Sheet 'Administered Calculation' NOT found!")
