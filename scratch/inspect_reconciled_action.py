import openpyxl

file_path = r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Infusion_Reconciliation_Report.xlsx"
wb = openpyxl.load_workbook(file_path)

if "Order Reconciliation Summary" in wb.sheetnames:
    ws = wb["Order Reconciliation Summary"]
    print("=== Columns ===")
    headers = [cell.value for cell in ws[1]]
    for idx, h in enumerate(headers, 1):
        print(f"  Col {idx}: {h}")
        
    print("\n=== Sample of First 5 rows of data ===")
    for r in range(2, 7):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)]
        print(f"  Row {r}: {row_vals}")
else:
    print("Sheet 'Order Reconciliation Summary' not found!")
