import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt
import os

def highlight_run(run):
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW

def update_document(file_path, output_path):
    if not os.path.exists(file_path):
        print(f"File not found: {file_path}")
        return

    doc = docx.Document(file_path)

    # 1. Add Header
    # Assuming 'Header' means adding a header to the document's header section
    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.text = "Infusion Waste Monitoring - Solution Design Document"
    for run in header_para.runs:
        run.bold = True

    # 2. Add Follow Up section at the beginning (after the title or executive summary, let's just insert it at the very top of the body)
    first_para = doc.paragraphs[0]
    follow_up_title = first_para.insert_paragraph_before("Follow Up:")
    follow_up_title.style = doc.styles['Heading 1'] if 'Heading 1' in doc.styles else doc.styles['Normal']
    
    follow_up_text = first_para.insert_paragraph_before("Still waiting to determine how PYXIS Data will be shared to Intelligent Automation email inbox, as there are current security restrictions preventing BOT Access. Currently, two factor authentication is being used for report access once sent to email")
    follow_up_text.style = doc.styles['Normal']
    
    # Highlight the newly added Follow Up text
    for run in follow_up_title.runs:
        highlight_run(run)
    for run in follow_up_text.runs:
        highlight_run(run)

    # Re-apply previous logic updates WITH highlighting
    for p in doc.paragraphs:
        # Update Reconciliation Logic
        if "Total Dispensed Amount =" in p.text:
            p.text = ""
            run = p.add_run("Total Dispensed Amount =\nSUM(Paired Vend Amounts)")
            highlight_run(run)
        elif "SUM(All Vend Amounts)" in p.text:
            p.text = p.text.replace("SUM(All Vend Amounts)", "")
            run = p.add_run("SUM(Paired Vend Amounts)")
            highlight_run(run)
            
        if "Total Waste Amount =" in p.text:
            p.text = ""
            run = p.add_run("Total Waste Amount =\nSUM(Paired Waste Amounts)")
            highlight_run(run)
        elif "SUM(All Waste Amounts)" in p.text:
            p.text = p.text.replace("SUM(All Waste Amounts)", "")
            run = p.add_run("SUM(Paired Waste Amounts)")
            highlight_run(run)

    # Insert Vend/Waste Pairing Logic
    for i, p in enumerate(doc.paragraphs):
        if p.text == "Date Matching Logic":
            new_p = p.insert_paragraph_before("Vend/Waste Pairing Logic")
            new_p.style = p.style
            for run in new_p.runs: highlight_run(run)
            if not new_p.runs: highlight_run(new_p.add_run("Vend/Waste Pairing Logic"))
            
            new_p2 = p.insert_paragraph_before("Unlike simple summation, Vends and Wastes must be chronologically paired. A Waste transaction is paired with a preceding Vend transaction that occurred within the maximum hang time window (e.g. 96 hours).")
            new_p2.style = doc.styles['Normal']
            for run in new_p2.runs: highlight_run(run)
            if not new_p2.runs: highlight_run(new_p2.add_run("Unlike simple summation, Vends and Wastes must be chronologically paired. A Waste transaction is paired with a preceding Vend transaction that occurred within the maximum hang time window (e.g. 96 hours)."))
            
            new_p3 = p.insert_paragraph_before("Transaction types 'Stock Return' and 'Vend (Cancel)' are excluded from active pairs. Any unmatched Wastes are flagged as exceptions.")
            new_p3.style = doc.styles['Normal']
            for run in new_p3.runs: highlight_run(run)
            if not new_p3.runs: highlight_run(new_p3.add_run("Transaction types 'Stock Return' and 'Vend (Cancel)' are excluded from active pairs. Any unmatched Wastes are flagged as exceptions."))
            break
            
    # Insert Controlled Substances Filtering
    for i, p in enumerate(doc.paragraphs):
        if "The documented waste amount is found in Quantity" in p.text:
            new_p = p.insert_paragraph_before("Controlled Substances Filtering:")
            new_p.style = p.style
            for run in new_p.runs: highlight_run(run)
            if not new_p.runs: highlight_run(new_p.add_run("Controlled Substances Filtering:"))
            
            new_p2 = p.insert_paragraph_before("The system exclusively filters for controlled substances (e.g., fentanyl, midazolam, versed, ketamine, morphine, hydromorphone, dilaudid) using the MedDescription field. Non-controlled substances are excluded from reconciliation.")
            new_p2.style = doc.styles['Normal']
            for run in new_p2.runs: highlight_run(run)
            if not new_p2.runs: highlight_run(new_p2.add_run("The system exclusively filters for controlled substances (e.g., fentanyl, midazolam, versed, ketamine, morphine, hydromorphone, dilaudid) using the MedDescription field. Non-controlled substances are excluded from reconciliation."))
            break

    doc.save(output_path)
    print(f"Updated document saved to {output_path}")

if __name__ == '__main__':
    input_doc = r"C:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Project Docs\SDD Infusion.docx"
    output_doc = r"C:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Project Docs\SDD Infusion_Updated.docx"
    update_document(input_doc, output_doc)
