import openpyxl

file_path = r"c:\Users\GaneshBhat\OneDrive - Novatio Solutions\Desktop\Infusion\Output\Step1_Pyxis_Analysis.xlsx"
wb = openpyxl.load_workbook(file_path)

if "Order Reconciliation Summary" in wb.sheetnames:
    ws = wb["Order Reconciliation Summary"]
    print("=== Order Reconciliation Summary ===")
    
    headers = [cell.value for cell in ws[1]]
    print("Headers (Length: {}):".format(len(headers)))
    for idx, h in enumerate(headers, 1):
        print(f"  Col {idx}: {h}")
        
    first_row = [cell.value for cell in ws[2]]
    print("\nFirst Row Data (Length: {}):".format(len(first_row)))
    for idx, (h, val) in enumerate(zip(headers, first_row), 1):
        print(f"  Col {idx} ({h}): {val}")
else:
    print("Sheet 'Order Reconciliation Summary' not found!")
